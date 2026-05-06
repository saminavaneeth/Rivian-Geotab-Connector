from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from datetime import datetime

# Create a new Document
doc = Document()

# Add title
title = doc.add_heading('Product Requirements Document', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

subtitle = doc.add_paragraph('Geotab OEM Connector Integration Platform')
subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
subtitle_format = subtitle.runs[0]
subtitle_format.font.size = Pt(14)
subtitle_format.font.italic = True

# Add metadata
metadata = doc.add_paragraph()
metadata.add_run('Date: ').bold = True
metadata.add_run(datetime.now().strftime("%B %d, %Y"))

metadata.add_run('\nVersion: ').bold = True
metadata.add_run('1.0')

metadata.add_run('\nStatus: ').bold = True
metadata.add_run('Draft')

doc.add_paragraph()

# 1. Executive Summary
doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph(
    'The OEM Connector Integration Platform aims to streamline vehicle telematics data collection and '
    'integration across multiple Original Equipment Manufacturers (OEMs). This platform will enable Geotab '
    'to expand market reach, reduce integration time-to-value, and enhance customer data accessibility while '
    'maintaining security and compliance standards.'
)

# 2. Problem Statement
doc.add_heading('2. Problem Statement', level=1)

doc.add_heading('2.1 Current Challenges', level=2)
challenges = [
    'Time-intensive OEM integration processes limit our ability to onboard new manufacturers quickly',
    'Disparate data formats and APIs across OEMs create maintenance overhead and increase error risk',
    'Customers lack unified visibility into vehicle data from multiple OEM sources',
    'Integration failures impact customer satisfaction and increase support costs'
]
for challenge in challenges:
    doc.add_paragraph(challenge, style='List Bullet')

doc.add_heading('2.2 Market Opportunity', level=2)
doc.add_paragraph(
    'The fleet management market is rapidly consolidating around unified telematics platforms. '
    'OEMs are increasingly opening APIs to enable third-party integrations. Organizations that can '
    'rapidly integrate new OEM data sources will capture greater market share and improve customer retention.'
)

# 3. Goals and Objectives
doc.add_heading('3. Goals and Objectives', level=1)

goals = [
    ('Reduce OEM Integration Time', 'Decrease time-to-market for new OEM connectors from 6+ months to 3 months'),
    ('Increase Data Coverage', 'Expand from current OEM coverage to support 10+ new manufacturer integrations'),
    ('Improve Data Quality', 'Achieve 99.9% data accuracy and uptime for all OEM connectors'),
    ('Enhance User Experience', 'Provide customers with a unified dashboard for multi-OEM fleet visibility'),
    ('Strengthen Compliance', 'Ensure all integrations meet GDPR, CCPA, and industry-specific regulations')
]

for goal, description in goals:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(goal + ': ').bold = True
    p.add_run(description)

# 4. Target Users and Personas
doc.add_heading('4. Target Users and Personas', level=1)

doc.add_heading('4.1 Primary Users', level=2)
personas = {
    'Fleet Managers': 'Need real-time visibility into vehicle performance across multiple brands',
    'Operations Teams': 'Require automated data pipelines and alerting across diverse OEM sources',
    'API Consumers': 'Demand standardized data schemas and REST endpoints'
}
for persona, need in personas.items():
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(persona + ': ').bold = True
    p.add_run(need)

# 5. Success Metrics
doc.add_heading('5. Success Metrics', level=1)

metrics_data = [
    ('Integration Time', '< 3 months per new OEM connector'),
    ('Data Availability', '99.9% uptime SLA'),
    ('Customer Adoption', '80% of enterprise customers using multi-OEM connectors'),
    ('Support Tickets', '30% reduction in integration-related support issues'),
    ('API Performance', '< 200ms median latency for data queries')
]

table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Metric'
hdr_cells[1].text = 'Target'

for metric, target in metrics_data:
    row_cells = table.add_row().cells
    row_cells[0].text = metric
    row_cells[1].text = target

# 6. Functional Requirements
doc.add_heading('6. Functional Requirements', level=1)

doc.add_heading('6.1 Connector Framework', level=2)
fr = [
    'Standardized connector architecture supporting OAuth 2.0 and API key authentication',
    'Automated data transformation engine converting OEM formats to Geotab schema',
    'Real-time and batch ingestion capabilities with configurable sync intervals',
    'Connector versioning and backward compatibility management'
]
for item in fr:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('6.2 Data Management', level=2)
dm = [
    'Unified vehicle data model consolidating information from multiple OEMs',
    'Data validation and reconciliation workflows',
    'Historical data retention with configurable archival policies',
    'Audit logging for compliance and troubleshooting'
]
for item in dm:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('6.3 API and Integration', level=2)
api = [
    'RESTful API with standardized endpoints for multi-OEM data queries',
    'GraphQL interface for complex data aggregation scenarios',
    'Webhook support for real-time event notifications',
    'SDK libraries for popular programming languages (Python, Java, Node.js)'
]
for item in api:
    doc.add_paragraph(item, style='List Bullet')

# 7. Non-Functional Requirements
doc.add_heading('7. Non-Functional Requirements', level=1)

nfr_table = doc.add_table(rows=1, cols=2)
nfr_table.style = 'Light Grid Accent 1'
nfr_hdr = nfr_table.rows[0].cells
nfr_hdr[0].text = 'Requirement'
nfr_hdr[1].text = 'Specification'

nfr_data = [
    ('Performance', 'API response time < 200ms (p95), support 10k+ req/min'),
    ('Reliability', '99.9% uptime SLA with automatic failover'),
    ('Security', 'End-to-end encryption, role-based access control, regular security audits'),
    ('Scalability', 'Horizontal scaling to support 100M+ vehicles'),
    ('Maintainability', 'Automated testing, CI/CD pipeline, comprehensive documentation')
]

for req, spec in nfr_data:
    row = nfr_table.add_row().cells
    row[0].text = req
    row[1].text = spec

# 8. Use Cases
doc.add_heading('8. Use Cases', level=1)

use_cases = [
    ('Multi-OEM Fleet Dashboard', 'A fleet manager views vehicle diagnostics and location data from Tesla, Ford, and GM vehicles in a unified dashboard'),
    ('Automated Compliance Reporting', 'Operations team receives weekly reports aggregating emissions and maintenance data across all OEM sources'),
    ('Proactive Maintenance Alerts', 'System automatically detects maintenance issues across OEM diagnostics and alerts operators before failures occur'),
    ('Custom Data Export', 'API consumer retrieves normalized vehicle data for 10k+ vehicles from 5 different OEMs for custom analytics')
]

for i, (use_case, description) in enumerate(use_cases, 1):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run('UC-{}: {}\n'.format(i, use_case)).bold = True
    p.add_run(description)

# 9. Dependencies and Constraints
doc.add_heading('9. Dependencies and Constraints', level=1)

doc.add_heading('9.1 Technical Dependencies', level=2)
deps = [
    'OEM API availability and stability',
    'Third-party authentication providers (OAuth providers)',
    'Cloud infrastructure scaling capabilities'
]
for dep in deps:
    doc.add_paragraph(dep, style='List Bullet')

doc.add_heading('9.2 Constraints', level=2)
constraints = [
    'OEM API rate limits and quota restrictions',
    'Regulatory compliance requirements (GDPR, CCPA, data residency)',
    'Legacy system integration complexity',
    'Resource allocation and timeline constraints'
]
for constraint in constraints:
    doc.add_paragraph(constraint, style='List Bullet')

# 10. Timeline and Milestones
doc.add_heading('10. Timeline and Milestones', level=1)

timeline_data = [
    ('Q2 2026', 'Architecture design, connector framework development'),
    ('Q3 2026', 'Pilot integrations with 3 major OEMs, beta testing'),
    ('Q4 2026', 'General availability, 8+ OEM connectors live'),
    ('Q1 2027', 'Advanced features (GraphQL, webhooks), performance optimization'),
    ('Q2 2027', 'Expansion to 15+ OEM partners, analytics enhancements')
]

timeline_table = doc.add_table(rows=1, cols=2)
timeline_table.style = 'Light Grid Accent 1'
timeline_hdr = timeline_table.rows[0].cells
timeline_hdr[0].text = 'Timeline'
timeline_hdr[1].text = 'Deliverables'

for period, deliverables in timeline_data:
    row = timeline_table.add_row().cells
    row[0].text = period
    row[1].text = deliverables

# 11. Risks and Mitigations
doc.add_heading('11. Risks and Mitigations', level=1)

risk_data = [
    ('OEM API Changes', 'Maintain abstraction layer, continuous monitoring, version management'),
    ('Data Security Breach', 'Implement encryption, regular audits, incident response plan'),
    ('Integration Delays', 'Dedicated OEM partnership team, early engagement, clear SLAs'),
    ('Adoption Challenges', 'User training programs, comprehensive documentation, customer support')
]

risk_table = doc.add_table(rows=1, cols=2)
risk_table.style = 'Light Grid Accent 1'
risk_hdr = risk_table.rows[0].cells
risk_hdr[0].text = 'Risk'
risk_hdr[1].text = 'Mitigation Strategy'

for risk, mitigation in risk_data:
    row = risk_table.add_row().cells
    row[0].text = risk
    row[1].text = mitigation

# 12. Success Criteria
doc.add_heading('12. Success Criteria', level=1)
success = [
    'All functional requirements implemented and tested',
    'Achieved 99.9% uptime during pilot period',
    'Successfully integrated with 8+ OEM APIs',
    'Customer feedback score >= 8/10',
    'Support team can handle integrations with < 1 hour ramp-up',
    'Zero critical security vulnerabilities in security audit'
]
for criterion in success:
    doc.add_paragraph(criterion, style='List Bullet')

# Save the document
doc.save('Geotab_OEM_Connector_PRD.docx')
print('PRD document created: Geotab_OEM_Connector_PRD.docx')
